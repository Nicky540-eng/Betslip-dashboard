"""Data layer for the Cashier Slip Dashboard.

Parses Slip Summary CSVs, stores them in Neon (Postgres) — or a local SQLite
file when no DATABASE_URL is set — and builds the Excel export.
"""
from __future__ import annotations

import hashlib
import io
from datetime import date, datetime

import pandas as pd
import sqlalchemy as sa

# ---------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------
SLIP_REQUIRED = ["Game", "Shop", "User", "Bet Slips", "First Slip Issued", "Last Slip Issued"]
MONEY_COLS = {
    "Paid In": "paid_in",
    "Withholding Tax": "withholding_tax",
    "Paid Out": "paid_out",
    "Winnings - Unpaid": "winnings_unpaid",
    "Net Win": "net_win",
}

metadata = sa.MetaData()

slips = sa.Table(
    "cs_slip_summary",
    metadata,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("month", sa.Date, nullable=False, index=True),
    sa.Column("shop", sa.String(100), nullable=False, index=True),
    sa.Column("cashier", sa.String(200), nullable=False),
    sa.Column("game", sa.String(200), nullable=False),
    sa.Column("bet_slips", sa.Integer, nullable=False),
    sa.Column("first_slip", sa.DateTime),
    sa.Column("last_slip", sa.DateTime),
    sa.Column("currency", sa.String(10)),
    *[sa.Column(c, sa.Numeric(14, 2)) for c in MONEY_COLS.values()],
    sa.Column("upload_id", sa.Integer, index=True),
    sa.UniqueConstraint("month", "shop", "cashier", "game", name="uq_cs_slip_row"),
)

uploads = sa.Table(
    "cs_uploads",
    metadata,
    sa.Column("id", sa.Integer, primary_key=True, autoincrement=True),
    sa.Column("file_name", sa.String(300)),
    sa.Column("file_md5", sa.String(32), nullable=False, unique=True),
    sa.Column("report_type", sa.String(50), nullable=False),
    sa.Column("month", sa.Date),
    sa.Column("row_count", sa.Integer),
    sa.Column("bet_slips", sa.Integer),
    sa.Column("uploaded_at", sa.DateTime),
)


def make_engine(url: str | None) -> sa.Engine:
    if not url:
        return sa.create_engine("sqlite:///local_dev.db")
    for prefix in ("postgres://", "postgresql://"):
        if url.startswith(prefix):
            # Pin the psycopg2 driver (SQLAlchemy 2.1 otherwise defaults to psycopg 3).
            url = "postgresql+psycopg2://" + url[len(prefix):]
    # pool_pre_ping: Neon suspends idle computes, so stale connections are common.
    return sa.create_engine(url, pool_pre_ping=True, pool_recycle=300)


def init_db(engine: sa.Engine) -> None:
    metadata.create_all(engine)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
class UploadError(Exception):
    pass


def md5_of(raw: bytes) -> str:
    return hashlib.md5(raw).hexdigest()


def _read_csv(raw: bytes) -> pd.DataFrame:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            df = pd.read_csv(io.BytesIO(raw), encoding=enc, dtype=str)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise UploadError("Could not read the file's text encoding.")
    df.columns = [str(c).strip().strip('"').strip() for c in df.columns]
    return df


def detect_report_type(raw: bytes) -> str:
    cols = set(_read_csv(raw).columns)
    if set(SLIP_REQUIRED) <= cols:
        return "slip_summary"
    if {"Shop", "User"} <= cols and "Game" not in cols:
        return "cash_operations"
    return "unknown"


def _to_number(s: pd.Series) -> pd.Series:
    cleaned = s.astype(str).str.replace(",", "", regex=False).str.replace(" ", "", regex=False)
    return pd.to_numeric(cleaned, errors="coerce")


def _to_datetime(s: pd.Series) -> pd.Series:
    s = s.astype(str).str.strip()
    # Report dates are DD/MM/YY — parse explicitly so 05/04/26 is 5 April, not 4 May.
    out = pd.to_datetime(s, format="%d/%m/%y %H:%M:%S", errors="coerce")
    if out.isna().mean() > 0.5:
        out = pd.to_datetime(s, dayfirst=True, errors="coerce")
    return out


def parse_slip_summary(raw: bytes) -> tuple[pd.DataFrame, date]:
    """Return clean rows (one per month × shop × cashier × game) and the report month."""
    df = _read_csv(raw)
    missing = [c for c in SLIP_REQUIRED if c not in df.columns]
    if missing:
        raise UploadError(f"Not a Slip Summary file — missing column(s): {', '.join(missing)}")

    df = df.dropna(subset=["Game", "Shop", "User"])
    for c in ("Game", "Shop", "User"):
        df[c] = df[c].str.strip()
    df = df[(df["Game"] != "") & (df["Shop"] != "") & (df["User"] != "")]
    df = df[~df["User"].str.lower().str.match(r"^(grand\s+)?totals?$")]
    df = df[~df["Game"].str.lower().str.match(r"^(grand\s+)?totals?$")]
    if df.empty:
        raise UploadError("The file has no cashier rows.")

    out = pd.DataFrame(
        {
            "shop": df["Shop"].values,
            "cashier": df["User"].values,
            "game": df["Game"].values,
            "bet_slips": _to_number(df["Bet Slips"]).fillna(0).astype(int).values,
            "first_slip": _to_datetime(df["First Slip Issued"]).values,
            "last_slip": _to_datetime(df["Last Slip Issued"]).values,
            "currency": (df["Currency"].str.strip() if "Currency" in df else pd.Series("ZAR", index=df.index)).values,
        }
    )
    for src, dst in MONEY_COLS.items():
        out[dst] = _to_number(df[src]).fillna(0).values if src in df else 0.0

    stamps = pd.concat([out["first_slip"], out["last_slip"]]).dropna()
    if stamps.empty:
        raise UploadError("Could not read any slip dates, so the month can't be determined.")
    periods = stamps.dt.to_period("M").unique()
    if len(periods) > 1:
        span = ", ".join(p.strftime("%b %Y") for p in sorted(periods))
        raise UploadError(f"File covers more than one month ({span}). Export one month per file.")
    month = periods[0].to_timestamp().date()

    # Collapse any repeated cashier×game rows so the unique key holds.
    out = (
        out.groupby(["shop", "cashier", "game"], as_index=False)
        .agg(
            bet_slips=("bet_slips", "sum"),
            first_slip=("first_slip", "min"),
            last_slip=("last_slip", "max"),
            currency=("currency", "first"),
            **{c: (c, "sum") for c in MONEY_COLS.values()},
        )
    )
    out.insert(0, "month", month)
    return out, month


# ---------------------------------------------------------------------------
# Writes
# ---------------------------------------------------------------------------
def find_md5(engine: sa.Engine, md5: str):
    with engine.connect() as conn:
        return conn.execute(
            sa.select(uploads.c.file_name, uploads.c.month, uploads.c.uploaded_at).where(uploads.c.file_md5 == md5)
        ).first()


def month_loaded(engine: sa.Engine, month: date) -> bool:
    with engine.connect() as conn:
        return (
            conn.execute(
                sa.select(sa.func.count()).select_from(slips).where(slips.c.month == month)
            ).scalar()
            > 0
        )


def _records(df: pd.DataFrame) -> list[dict]:
    recs = []
    for row in df.to_dict("records"):
        clean = {}
        for k, v in row.items():
            if v is None or (not isinstance(v, (str, date)) and pd.isna(v)):
                clean[k] = None
            elif isinstance(v, pd.Timestamp):
                clean[k] = v.to_pydatetime()
            elif hasattr(v, "item"):  # numpy scalar -> python
                clean[k] = v.item()
            else:
                clean[k] = v
        recs.append(clean)
    return recs


def save_slip_summary(engine: sa.Engine, file_name: str, md5: str, df: pd.DataFrame, month: date, replace: bool) -> str:
    with engine.begin() as conn:
        exists = conn.execute(
            sa.select(sa.func.count()).select_from(slips).where(slips.c.month == month)
        ).scalar()
        if exists and not replace:
            raise UploadError(f"{month:%B %Y} is already loaded. Tick 'Replace' to overwrite it.")
        if exists:
            conn.execute(slips.delete().where(slips.c.month == month))
            conn.execute(
                uploads.delete().where((uploads.c.month == month) & (uploads.c.report_type == "slip_summary"))
            )
        upload_id = conn.execute(
            uploads.insert().values(
                file_name=file_name,
                file_md5=md5,
                report_type="slip_summary",
                month=month,
                row_count=len(df),
                bet_slips=int(df["bet_slips"].sum()),
                uploaded_at=datetime.utcnow(),
            )
        ).inserted_primary_key[0]
        conn.execute(slips.insert(), _records(df.assign(upload_id=upload_id)))
    return "replaced" if exists else "added"


def delete_month(engine: sa.Engine, month: date) -> None:
    with engine.begin() as conn:
        conn.execute(slips.delete().where(slips.c.month == month))
        conn.execute(uploads.delete().where(uploads.c.month == month))


# ---------------------------------------------------------------------------
# Reads
# ---------------------------------------------------------------------------
def data_version(engine: sa.Engine) -> str:
    with engine.connect() as conn:
        n, mx = conn.execute(sa.select(sa.func.count(), sa.func.max(uploads.c.id))).first()
    return f"{n}-{mx}"


def load_slips(engine: sa.Engine) -> pd.DataFrame:
    q = sa.select(
        slips.c.month, slips.c.shop, slips.c.cashier, slips.c.game, slips.c.bet_slips,
        slips.c.paid_in, slips.c.paid_out, slips.c.net_win,
    )
    df = pd.read_sql(q, engine)
    df["month"] = pd.to_datetime(df["month"])
    df["bet_slips"] = df["bet_slips"].astype(int)
    for c in ("paid_in", "paid_out", "net_win"):
        df[c] = df[c].astype(float)
    return df


def load_upload_log(engine: sa.Engine) -> pd.DataFrame:
    q = sa.select(
        uploads.c.month, uploads.c.file_name, uploads.c.row_count, uploads.c.bet_slips, uploads.c.uploaded_at
    ).order_by(uploads.c.month)
    return pd.read_sql(q, engine)


# ---------------------------------------------------------------------------
# Selected month + all-months total (used by the page and the Excel export)
# ---------------------------------------------------------------------------
def month_and_total(df: pd.DataFrame, rows: list[str], month, month_label: str, total_label: str) -> pd.DataFrame:
    total = df.groupby(rows)["bet_slips"].sum().rename(total_label)
    in_month = df[df["month"] == pd.Timestamp(month)].groupby(rows)["bet_slips"].sum().rename(month_label)
    out = pd.concat([in_month, total], axis=1).fillna(0).astype(int)
    return out.sort_values([month_label, total_label], ascending=False).reset_index()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def build_excel(sheets: dict[str, pd.DataFrame]) -> bytes:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(frame.columns, start=1):
                width = max([len(col)] + [len(f"{v:,}" if isinstance(v, int) else str(v)) for v in frame[col]])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 3, 40)
                if pd.api.types.is_integer_dtype(frame[col]):
                    for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        cell.number_format = "#,##0"
    return buf.getvalue()


# ---------------------------------------------------------------------------
# One-time seed of the historical months bundled in seed_data/
# ---------------------------------------------------------------------------
def seed_from_folder(engine: sa.Engine, folder: str) -> list[str]:
    """Load any bundled CSV whose month isn't in the database yet. Safe to run on every start."""
    import os

    loaded = []
    if not os.path.isdir(folder):
        return loaded
    for name in sorted(os.listdir(folder)):
        if not name.lower().endswith(".csv"):
            continue
        raw = open(os.path.join(folder, name), "rb").read()
        md5 = md5_of(raw)
        if find_md5(engine, md5):
            continue
        df, month = parse_slip_summary(raw)
        if month_loaded(engine, month):
            continue
        save_slip_summary(engine, name, md5, df, month, replace=False)
        loaded.append(f"{month:%b %Y}")
    return loaded
